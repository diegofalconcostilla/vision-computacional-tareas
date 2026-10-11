"""Versiones editables en Word (.docx) de los informes, generadas con pandoc.

- Informes de análisis: `Reporte.guardar()` ya escribe el .docx junto al PDF (vía `html_a_docx`).
- Prácticas: se convierten desde el notebook ejecutado en `codigo/`, así las ecuaciones
  quedan como ecuaciones de Word y el código como texto editable (`notebook_a_docx`).

Regenerar los .docx de todas las prácticas sin volver a ejecutar nada:

    python -m herramientas.word

Requiere pandoc (https://pandoc.org) y python-docx.
"""
import copy
import shutil
import subprocess
import tempfile
from pathlib import Path

import nbformat
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt, RGBColor

from .reporte import CURSO, EQUIPO, FECHA, INTEGRANTES

PLANTILLA = Path(__file__).with_name("plantilla_word.docx")
PLANTILLA_INFORME = Path(__file__).with_name("plantilla_informe.docx")  # informes "Relaciona lo aprendido"
TITULOS = RGBColor(0, 0, 0)  # títulos y subtítulos en negro
RAIZ = Path(__file__).resolve().parent.parent

def equipo(fecha):
    """Datos de portada de una entrega en equipo: integrantes con matrícula y fecha de entrega de Canvas."""
    return {"integrantes": EQUIPO, "fecha": fecha}


ENTREGA_EQUIPO = equipo("27/9/2026")
# Prácticas: carpeta -> (notebooks en orden, nombre del .docx, título, subtítulo, datos de portada). Mismos textos que los PDF.
PRACTICAS = {
    "04-practica-1-procesamiento-basico": (["practica_1_procesamiento_basico"], "Practica1_Procesamiento_basico",
        "Práctica 1. Procesamiento básico de imágenes", "Transformaciones píxel a píxel: fotométricas, negativo, gamma y sustracción",
        ENTREGA_EQUIPO),
    "05-practica-2-algoritmos-de-mejoramiento-de-imagenes-basado-por-pixeles": (
        ["practica_2_algoritmos_de_mejoramiento_de_imagenes_basado_por_pixeles"], "Practica2_Algoritmos_de_mejoramiento_de_imagenes_basado_por_pixeles",
        "Práctica 2. Algoritmos de mejoramiento de imágenes basado por pixeles", "Mosaicos, ventana deslizante (SWAHE) y CLAHE",
        ENTREGA_EQUIPO),
    "06-practica-3-filtros-espaciales": (["practica_3_filtros_espaciales"], "Practica3_Filtros_espaciales",
        "Práctica 3. Algoritmos de mejoramiento de imágenes en el dominio espacial", "Prewitt, Sobel y Laplaciano sin OpenCV, realce por diferencia y mejora de una gammagrafía ósea",
        equipo("4/10/2026")),
    "08-practica-4-morfologia": (["practica_4_morfologia"], "Practica4_Morfologia",
        "Práctica 4. Algoritmos de mejoramiento de imágenes usando operaciones morfológicas", "Erosión, dilatación, apertura y cierre aplicados y medidos",
        equipo("11/10/2026")),
    "10-practica-5-fourier": (["practica_5_fourier"], "Practica5_Fourier",
        "Práctica 5. Algoritmos de mejoramiento de imágenes en el dominio de Fourier", "Filtros pasa bajas y pasa altas: ideal, Butterworth y gaussiano",
        equipo("18/10/2026")),
    "12-practica-6-canny": (["practica_6_canny"], "Practica6_Canny",
        "Práctica 6. Extracción de líneas con algoritmo de Canny", "Efecto de σ, del suavizado previo y del tipo de imagen",
        equipo("25/10/2026")),
    "14-practica-7-harris-matching": (["practica_7a_harris", "practica_7b_matching"], "Practica7_Harris_y_matching",
        "Práctica 7. Algoritmos de extracción de características",
        "Parte A: Harris bajo cambios de iluminación, ángulo y escala · Parte B: Harris, ORB y SIFT en matching",
        equipo("1/11/2026")),
    "16-practica-8-otsu": (["practica_8_otsu"], "Practica8_Otsu",
        "Práctica 8. Algoritmo Otsu para segmentación", "Umbral único frente a Otsu, imágenes de estilos distintos y Otsu por ventanas",
        equipo("8/11/2026")),
}


# ---------------------------------------------------------------- plantilla de estilos

def plantilla():
    """Crea (una vez) la plantilla de estilos: A4, Calibri 10.5 pt y títulos en negro, como en los PDF."""
    if PLANTILLA.exists():
        return PLANTILLA
    datos = subprocess.run(["pandoc", "--print-default-data-file", "reference.docx"],
                           capture_output=True, check=True).stdout
    PLANTILLA.write_bytes(datos)
    doc = Document(PLANTILLA)
    for s in doc.sections:
        s.page_width, s.page_height = Mm(210), Mm(297)
        s.left_margin = s.right_margin = Mm(16)
        s.top_margin = s.bottom_margin = Mm(17)
        s.header_distance = s.footer_distance = Mm(10)  # w:pgMar exige los siete atributos
        s.gutter = Mm(0)
    # doc.styles["Heading 1"] no sirve aquí: python-docx lo traduce a "heading 1" y pandoc lo nombra con mayúscula
    est = {s.name: s for s in doc.styles}
    for nombre in ("Normal", "Body Text", "First Paragraph", "Compact"):
        f = est[nombre].font
        f.name, f.size = "Calibri", Pt(10.5)
    for nombre, tam in (("Title", 22), ("Subtitle", 13), ("Heading 1", 18), ("Heading 2", 14), ("Heading 3", 11.5), ("Heading 4", 10.5)):
        f = est[nombre].font
        f.name, f.size, f.bold = "Calibri", Pt(tam), nombre != "Subtitle"
        f.color.rgb = TITULOS
    est["Verbatim Char"].font.name, est["Verbatim Char"].font.size = "Consolas", Pt(8.5)
    doc.save(PLANTILLA)
    return PLANTILLA


def plantilla_informe():
    """Crea (una vez) la plantilla de los informes: A4, Aptos 11 pt justificado y encabezados sin negrita en negro,
    como el informe de ejemplo hecho en Word."""
    if PLANTILLA_INFORME.exists():
        return PLANTILLA_INFORME
    datos = subprocess.run(["pandoc", "--print-default-data-file", "reference.docx"],
                           capture_output=True, check=True).stdout
    PLANTILLA_INFORME.write_bytes(datos)
    doc = Document(PLANTILLA_INFORME)
    for s in doc.sections:
        s.page_width, s.page_height = Mm(210), Mm(297)
        s.left_margin = s.right_margin = Mm(25)
        s.top_margin, s.bottom_margin = Mm(25), Mm(22)
        s.header_distance = s.footer_distance = Mm(12)
        s.gutter = Mm(0)
    est = {s.name: s for s in doc.styles}
    for nombre in ("Normal", "Body Text", "First Paragraph", "Compact"):
        e = est[nombre]
        e.font.name, e.font.size, e.font.color.rgb = "Aptos", Pt(11), RGBColor(0, 0, 0)
        if nombre != "Normal":
            e.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            e.paragraph_format.line_spacing = 1.15
            e.paragraph_format.space_after = Pt(8)
    for nombre, tam, color in (("Title", 27, RGBColor(0, 0, 0)), ("Heading 1", 27, RGBColor(0, 0, 0)),
                               ("Heading 2", 17, TITULOS), ("Heading 3", 13.5, TITULOS),
                               ("Heading 4", 12, TITULOS), ("Subtitle", 14, TITULOS)):
        f = est[nombre].font
        f.name, f.size, f.bold, f.italic, f.color.rgb = "Aptos Display", Pt(tam), False, False, color
    est["Subtitle"].font.name = est["Heading 4"].font.name = "Aptos"
    for nombre in ("Subtitle", "Heading 4"):  # letras espaciadas, como en el ejemplo
        rpr = est[nombre].element.get_or_add_rPr()
        sp = OxmlElement("w:spacing")
        sp.set(qn("w:val"), "20")
        rpr.append(sp)
    doc.save(PLANTILLA_INFORME)
    return PLANTILLA_INFORME


def _pandoc(entrada, salida, formato, recursos, referencia=None):
    if shutil.which("pandoc") is None:
        raise RuntimeError("Se necesita pandoc para generar el .docx (https://pandoc.org/installing.html)")
    subprocess.run(["pandoc", str(entrada), "-f", formato, "-t", "docx", "-o", str(salida),
                    f"--reference-doc={referencia or plantilla()}", f"--resource-path={recursos}"], check=True)


# ---------------------------------------------------------------- retoques con python-docx

# Orden que exige el esquema OOXML para los hijos de w:tblPr y w:tcPr (Word rechaza otros órdenes)
_ORDEN_TBLPR = ("tblStyle", "tblpPr", "tblOverlap", "bidiVisual", "tblStyleRowBandSize", "tblStyleColBandSize",
                "tblW", "jc", "tblCellSpacing", "tblInd", "tblBorders", "shd", "tblLayout", "tblCellMar", "tblLook",
                "tblCaption", "tblDescription")
_ORDEN_TCPR = ("cnfStyle", "tcW", "gridSpan", "hMerge", "vMerge", "tcBorders", "shd", "noWrap", "tcMar",
               "textDirection", "tcFitText", "vAlign", "hideMark")


def _insertar_en_orden(padre, hijo, orden):
    nombre = hijo.tag.split("}")[1]
    for viejo in padre.findall(qn(f"w:{nombre}")):
        padre.remove(viejo)
    despues = set(orden[orden.index(nombre) + 1:])
    for i, h in enumerate(padre):
        if h.tag.split("}")[1] in despues:
            padre.insert(i, hijo)
            return
    padre.append(hijo)


def _bordes_tabla(tabla):
    bordes = OxmlElement("w:tblBorders")
    for lado in ("top", "left", "bottom", "right", "insideH", "insideV"):
        b = OxmlElement(f"w:{lado}")
        b.set(qn("w:val"), "single"), b.set(qn("w:sz"), "4"), b.set(qn("w:color"), "C9D3EA")
        bordes.append(b)
    _insertar_en_orden(tabla._tbl.tblPr, bordes, _ORDEN_TBLPR)
    for celda in tabla.rows[0].cells:  # encabezado azul con texto blanco, como en los PDF
        sh = OxmlElement("w:shd")
        sh.set(qn("w:val"), "clear"), sh.set(qn("w:color"), "auto"), sh.set(qn("w:fill"), "0B3D91")
        _insertar_en_orden(celda._tc.get_or_add_tcPr(), sh, _ORDEN_TCPR)
        for par in celda.paragraphs:
            for run in par.runs:
                run.font.color.rgb, run.font.bold = RGBColor(0xFF, 0xFF, 0xFF), True


def _salto_antes(parrafo):
    parrafo.insert_paragraph_before().add_run().add_break(WD_BREAK.PAGE)


def _pulir(ruta_docx, portada=None):
    """Bordes en tablas, portada (si se da) y salto de página después de la portada."""
    doc = Document(ruta_docx)
    for t in doc.tables:
        _bordes_tabla(t)
    cuerpo = doc.paragraphs
    if portada:
        estilos = {s.name: s for s in doc.styles}
        primero = cuerpo[0]
        for i, (texto, estilo) in enumerate(portada):
            p = primero.insert_paragraph_before(texto, style=estilos[estilo])
            if i == 0:  # nombre del curso en gris, como en los PDF
                p.runs[0].font.color.rgb = RGBColor(0x55, 0x55, 0x55)
        _salto_antes(primero)
    else:
        # informes HTML: título de la actividad (Title) y tema (Subtitle), seguidos sin salto de la sección "Equipo"
        estilos = {s.name: s for s in doc.styles}
        h1 = next((i for i, p in enumerate(cuerpo) if p.style.name == "Heading 1"), None)
        if h1 is not None:
            for p in cuerpo[:h1]:  # pandoc repite el <title> del HTML como bloque de título
                if p.style.name == "Title":
                    p._element.getparent().remove(p._element)
            cuerpo = doc.paragraphs
            h1 = next(i for i, p in enumerate(cuerpo) if p.style.name == "Heading 1")
            cuerpo[h1].style = estilos["Title"]
            sig = cuerpo[h1 + 1] if h1 + 1 < len(cuerpo) else None
            if sig is not None and not sig.style.name.startswith("Heading"):
                sig.style = estilos["Subtitle"]
    doc.save(ruta_docx)
    return ruta_docx


# ---------------------------------------------------------------- API

def html_a_docx(ruta_html, ruta_docx, recursos=None):
    """Convierte el HTML de un informe (el mismo que se imprime a PDF) a .docx."""
    ruta_html = Path(ruta_html)
    _pandoc(ruta_html, ruta_docx, "html", recursos or ruta_html.parent, plantilla_informe())
    return _pulir(ruta_docx)


def notebook_a_docx(notebooks, ruta_docx, titulo, subtitulo="", integrantes=None,
                    entrega="Notebook ejecutado + informe de resultados", fecha=FECHA):
    """Une uno o más notebooks ejecutados en un solo .docx con portada."""
    notebooks = [n if hasattr(n, "cells") else nbformat.read(str(n), as_version=4) for n in notebooks]
    unido = copy.deepcopy(notebooks[0])
    for nb in notebooks[1:]:
        unido.cells.extend(copy.deepcopy(nb.cells))
    for c in unido.cells:  # quita el texto "<Figure size …>" que acompaña a cada gráfica
        if c.cell_type == "code":
            for o in c.get("outputs", []):
                if "image/png" in o.get("data", {}):
                    o["data"].pop("text/plain", None)
    portada = ([(CURSO, "Body Text"), (titulo, "Title"), (subtitulo, "Subtitle"), ("Integrantes:", "Body Text")]
               + [(i, "Body Text") for i in (integrantes or INTEGRANTES)]  # uno por línea, como en el PDF
               + [(f"Fecha: {fecha}", "Body Text"), (f"Entrega: {entrega}", "Body Text")])
    with tempfile.TemporaryDirectory() as tmp:
        ruta_nb = Path(tmp) / "unido.ipynb"
        nbformat.write(unido, str(ruta_nb))
        _pandoc(ruta_nb, ruta_docx, "ipynb", tmp)
    return _pulir(ruta_docx, portada=[x for x in portada if x[0]])


def practicas():
    """Genera el .docx de cada práctica a partir de sus notebooks ya ejecutados."""
    hechos = []
    for carpeta, (nbs, nombre, titulo, sub, *extra) in PRACTICAS.items():
        base = RAIZ / "actividades" / carpeta
        rutas = [base / "codigo" / f"{n}.ipynb" for n in nbs]
        hechos.append(notebook_a_docx(rutas, base / f"{nombre}.docx", titulo, sub, **(extra[0] if extra else {})))
        print("DOCX", hechos[-1].relative_to(RAIZ))
    return hechos


if __name__ == "__main__":
    practicas()
